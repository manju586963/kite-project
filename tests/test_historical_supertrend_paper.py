#!/usr/bin/env python3
"""Unit tests for intraday Supertrend paper trading (no network)."""

from __future__ import annotations

import unittest
from datetime import datetime, timedelta
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from historical_supertrend_paper import (  # noqa: E402
    Candle,
    CandleST,
    allows_new_entry,
    compute_supertrend,
    filter_last_trading_days,
    is_square_off_candle,
    make_output_dir,
    simulate_paper_trades,
    summarize_trades,
    to_heikin_ashi,
    unique_file_path,
    write_monthly_excel,
    write_consolidated_monthly_summary,
)


def _candle(ts: datetime, close: float = 100.0) -> CandleST:
    return CandleST(
        date=ts,
        open=close,
        high=close + 1,
        low=close - 1,
        close=close,
        volume=10,
        signal=None,
        direction="bullish",
        supertrend=close - 5,
        atr=1.0,
    )


class IntradayRuleTests(unittest.TestCase):
    def test_heikin_ashi_conversion(self) -> None:
        raw = [
            Candle(datetime(2026, 8, 1, 10, 0), 100, 110, 90, 105),
            Candle(datetime(2026, 8, 1, 10, 15), 105, 120, 100, 115),
        ]
        ha = to_heikin_ashi(raw)
        self.assertEqual(len(ha), 2)
        # First HA close = (100+110+90+105)/4 = 101.25
        self.assertAlmostEqual(ha[0].close, 101.25)
        # First HA open = (100+105)/2 = 102.5
        self.assertAlmostEqual(ha[0].open, 102.5)
        self.assertAlmostEqual(ha[0].high, max(110, 102.5, 101.25))
        self.assertAlmostEqual(ha[0].low, min(90, 102.5, 101.25))
        # Second HA open = (prev_ha_open + prev_ha_close) / 2
        self.assertAlmostEqual(ha[1].open, (ha[0].open + ha[0].close) / 2)
        self.assertAlmostEqual(ha[1].close, (105 + 120 + 100 + 115) / 4)

    def test_supertrend_runs_on_heikin_ashi(self) -> None:
        start = datetime(2026, 7, 1, 10, 0)
        prices = [(100, 112, 95, 108)] + [
            (108 + i, 110 + i, 106 + i, 109 + i) for i in range(40)
        ]
        raw = [
            Candle(start + timedelta(minutes=15 * i), o, h, l, c)
            for i, (o, h, l, c) in enumerate(prices)
        ]
        ha = to_heikin_ashi(raw)
        rows = compute_supertrend(ha, period=10, multiplier=1)
        self.assertEqual(len(ha), len(raw))
        self.assertTrue(any(r.direction for r in rows))
        self.assertAlmostEqual(ha[0].close, (100 + 112 + 95 + 108) / 4)
        self.assertAlmostEqual(ha[0].open, (100 + 108) / 2)
        self.assertAlmostEqual(ha[1].open, (ha[0].open + ha[0].close) / 2)

    def test_entry_cutoff_and_square_off_times(self) -> None:
        self.assertTrue(allows_new_entry(datetime(2026, 8, 1, 22, 45)))
        self.assertTrue(allows_new_entry(datetime(2026, 8, 1, 23, 0)))
        self.assertFalse(allows_new_entry(datetime(2026, 8, 1, 23, 15)))
        self.assertFalse(is_square_off_candle(datetime(2026, 8, 1, 23, 0)))
        self.assertTrue(is_square_off_candle(datetime(2026, 8, 1, 23, 15)))

    def test_supertrend_10_1_emits_flips(self) -> None:
        start = datetime(2026, 7, 1, 10, 0)
        prices = (
            [(100 - i, 101 - i, 99 - i, 100 - i) for i in range(20)]
            + [(80 + i, 81 + i, 79 + i, 80 + i) for i in range(20)]
            + [(100 - i, 101 - i, 99 - i, 100 - i) for i in range(20)]
        )
        candles = [
            Candle(
                date=start + timedelta(minutes=15 * i),
                open=o,
                high=h,
                low=l,
                close=c,
            )
            for i, (o, h, l, c) in enumerate(prices)
        ]
        rows = compute_supertrend(candles, period=10, multiplier=1)
        signals = [r.signal for r in rows if r.signal]
        self.assertIn("BUY", signals)
        self.assertIn("SELL", signals)

    def test_intraday_square_off_at_2315(self) -> None:
        d = datetime(2026, 8, 1)
        rows = [
            _candle(d.replace(hour=10, minute=0), 100),
            _candle(d.replace(hour=11, minute=0), 105),
            _candle(d.replace(hour=23, minute=15), 102),
        ]
        rows[0].signal = "BUY"
        rows[1].signal = None
        rows[2].signal = "SELL"  # after cutoff / at square-off — ignored for entry

        trades = simulate_paper_trades(rows, lot_size=1)
        self.assertEqual(len(trades), 1)
        self.assertEqual(trades[0].side, "BUY")
        self.assertEqual(trades[0].exit_reason, "INTRADAY_SQUARE_OFF")
        self.assertEqual(trades[0].exit_price, 102)
        self.assertAlmostEqual(trades[0].points or 0, 2)

    def test_no_new_entry_after_2300(self) -> None:
        d = datetime(2026, 8, 1)
        rows = [
            _candle(d.replace(hour=22, minute=30), 100),
            _candle(d.replace(hour=23, minute=0), 101),  # last allowed entry time
            _candle(d.replace(hour=23, minute=15), 99),
        ]
        rows[0].signal = None
        rows[1].signal = "BUY"
        rows[2].signal = "SELL"

        trades = simulate_paper_trades(rows, lot_size=1)
        self.assertEqual(len(trades), 1)
        self.assertEqual(trades[0].entry_time.hour, 23)
        self.assertEqual(trades[0].entry_time.minute, 0)
        self.assertEqual(trades[0].exit_reason, "INTRADAY_SQUARE_OFF")

    def test_daily_reset_does_not_carry_position(self) -> None:
        day1 = datetime(2026, 8, 1)
        day2 = datetime(2026, 8, 2)
        rows = [
            _candle(day1.replace(hour=10, minute=0), 100),
            _candle(day1.replace(hour=23, minute=15), 110),
            _candle(day2.replace(hour=10, minute=0), 111),
            _candle(day2.replace(hour=23, minute=15), 105),
        ]
        rows[0].signal = "BUY"
        rows[1].signal = None
        rows[2].signal = None  # no fresh signal on day 2 → stays flat
        rows[3].signal = None

        trades = simulate_paper_trades(rows, lot_size=1)
        self.assertEqual(len(trades), 1)
        self.assertEqual(trades[0].trade_date.isoformat(), "2026-08-01")
        self.assertEqual(trades[0].exit_reason, "INTRADAY_SQUARE_OFF")

    def test_flip_within_day(self) -> None:
        d = datetime(2026, 8, 1)
        rows = [
            _candle(d.replace(hour=10, minute=0), 100),
            _candle(d.replace(hour=12, minute=0), 110),
            _candle(d.replace(hour=23, minute=15), 108),
        ]
        rows[0].signal = "BUY"
        rows[1].signal = "SELL"
        rows[2].signal = None

        trades = simulate_paper_trades(rows, lot_size=10)
        self.assertEqual(len(trades), 2)
        self.assertEqual(trades[0].exit_reason, "FLIP_SELL")
        self.assertEqual(trades[1].side, "SELL")
        self.assertEqual(trades[1].exit_reason, "INTRADAY_SQUARE_OFF")
        self.assertAlmostEqual(trades[0].pnl or 0, 100.0)  # 10 pts * lot 10

    def test_filter_last_trading_days(self) -> None:
        rows = []
        for day in range(1, 6):
            rows.append(_candle(datetime(2026, 8, day, 10, 0)))
        filtered = filter_last_trading_days(rows, trading_days=2)
        days = sorted({r.date.date() for r in filtered})
        self.assertEqual(days, [datetime(2026, 8, 4).date(), datetime(2026, 8, 5).date()])

    def test_summary_includes_square_offs(self) -> None:
        d = datetime(2026, 8, 1)
        rows = [
            _candle(d.replace(hour=10, minute=0), 100),
            _candle(d.replace(hour=23, minute=15), 101),
        ]
        rows[0].signal = "BUY"
        trades = simulate_paper_trades(rows, lot_size=1)
        stats = summarize_trades(trades)
        self.assertEqual(stats["square_offs"], 1)
        self.assertEqual(stats["total_trades"], 1)
        self.assertEqual(stats["max_loss_stops"], 0)

    def test_max_loss_stop_long_squares_off(self) -> None:
        d = datetime(2026, 8, 1)
        rows = [
            _candle(d.replace(hour=10, minute=0), 6000),
            _candle(d.replace(hour=11, minute=0), 5900),  # will set low below stop
            _candle(d.replace(hour=23, minute=15), 5800),
        ]
        rows[0].signal = "BUY"
        rows[1].signal = None
        rows[1].low = 6000 - 130  # breaches 125-point stop
        rows[1].high = 6000 - 10
        rows[2].signal = None

        trades = simulate_paper_trades(rows, lot_size=2, max_loss_points=125)
        self.assertEqual(len(trades), 1)
        self.assertEqual(trades[0].exit_reason, "MAX_LOSS_STOP")
        self.assertEqual(trades[0].stop_price, 5875)
        self.assertEqual(trades[0].exit_price, 5875)
        self.assertAlmostEqual(trades[0].points or 0, -125)
        self.assertAlmostEqual(trades[0].pnl or 0, -250)  # 125 × lot 2

    def test_max_loss_stop_before_reversal_then_reentry(self) -> None:
        d = datetime(2026, 8, 1)
        rows = [
            _candle(d.replace(hour=10, minute=0), 6000),
            _candle(d.replace(hour=11, minute=0), 5850),
            _candle(d.replace(hour=23, minute=15), 5860),
        ]
        rows[0].signal = "BUY"
        rows[1].signal = "SELL"  # reversal at close, but stop hits first on OHLC
        rows[1].low = 5800
        rows[1].high = 5900
        rows[1].close = 5850
        rows[2].signal = None

        trades = simulate_paper_trades(rows, lot_size=1, max_loss_points=125)
        # Stop closes LONG at 5875, then SELL opens SHORT at 5850 close.
        self.assertEqual(len(trades), 2)
        self.assertEqual(trades[0].exit_reason, "MAX_LOSS_STOP")
        self.assertEqual(trades[0].exit_price, 5875)
        self.assertEqual(trades[1].side, "SELL")
        self.assertEqual(trades[1].entry_price, 5850)
        self.assertEqual(trades[1].exit_reason, "INTRADAY_SQUARE_OFF")

    def test_max_loss_stop_short(self) -> None:
        d = datetime(2026, 8, 1)
        rows = [
            _candle(d.replace(hour=10, minute=0), 6000),
            _candle(d.replace(hour=12, minute=0), 6150),
            _candle(d.replace(hour=23, minute=15), 6160),
        ]
        rows[0].signal = "SELL"
        rows[1].signal = None
        rows[1].high = 6000 + 140
        rows[1].low = 6010
        rows[2].signal = None

        trades = simulate_paper_trades(rows, lot_size=1, max_loss_points=125)
        self.assertEqual(trades[0].exit_reason, "MAX_LOSS_STOP")
        self.assertEqual(trades[0].exit_price, 6125)
        self.assertAlmostEqual(trades[0].points or 0, -125)

    def test_output_dir_is_dated_and_unique(self) -> None:
        import shutil
        from historical_supertrend_paper import OUTPUT_ROOT

        fixed = datetime(2026, 8, 1, 16, 41, 5)
        first = make_output_dir(fixed)
        second = make_output_dir(fixed)
        try:
            self.assertEqual(first.name, "01-aug-2026")
            self.assertEqual(second.name, "01-aug-2026_2")
            self.assertNotEqual(first, second)
            self.assertTrue(first.is_dir())
            self.assertTrue(second.is_dir())
            self.assertEqual(first.parent, OUTPUT_ROOT)
        finally:
            shutil.rmtree(first, ignore_errors=True)
            shutil.rmtree(second, ignore_errors=True)

    def test_unique_file_path_never_overwrites(self) -> None:
        import shutil
        from historical_supertrend_paper import OUTPUT_ROOT

        folder = make_output_dir(datetime(2026, 8, 2, 10, 0, 0))
        try:
            first = unique_file_path(folder, "supertrend_signals.csv")
            first.write_text("a", encoding="utf-8")
            second = unique_file_path(folder, "supertrend_signals.csv")
            second.write_text("b", encoding="utf-8")
            self.assertEqual(first.name, "supertrend_signals.csv")
            self.assertEqual(second.name, "supertrend_signals_2.csv")
            self.assertTrue(first.exists())
            self.assertTrue(second.exists())
            self.assertEqual(first.read_text(encoding="utf-8"), "a")
            self.assertEqual(second.read_text(encoding="utf-8"), "b")
        finally:
            shutil.rmtree(folder, ignore_errors=True)
            leftover = OUTPUT_ROOT / "02-aug-2026"
            if leftover.exists() and not any(leftover.iterdir()):
                leftover.rmdir()

    def test_monthly_excel_export(self) -> None:
        import shutil
        from openpyxl import load_workbook
        from historical_supertrend_paper import OUTPUT_ROOT

        d1 = datetime(2026, 6, 1, 10, 0)
        d2 = datetime(2026, 7, 1, 10, 0)
        rows = [
            _candle(d1, 100),
            _candle(d1.replace(hour=23, minute=15), 110),
            _candle(d2, 200),
            _candle(d2.replace(hour=23, minute=15), 190),
        ]
        rows[0].signal = "BUY"
        rows[1].signal = None
        rows[2].signal = "SELL"
        rows[3].signal = None
        trades = simulate_paper_trades(rows, lot_size=1, max_loss_points=125)

        folder = make_output_dir(datetime(2026, 8, 3, 12, 0, 0))
        try:
            files = write_monthly_excel(
                folder,
                rows,
                trades,
                symbol="CRUDEOIL26AUGFUT",
                lot_size=1,
                max_loss_points=125,
            )
            names = sorted(p.name for p in files)
            self.assertEqual(names, ["jul-2026.xlsx", "jun-2026.xlsx"])
            wb = load_workbook(files[0] if files[0].name.startswith("jun") else files[1])
            # find jun file
            jun = next(p for p in files if p.name.startswith("jun"))
            wb = load_workbook(jun)
            self.assertEqual(set(wb.sheetnames), {"signals", "trades", "summary"})
            self.assertGreater(wb["signals"].max_row, 1)
            self.assertGreater(wb["trades"].max_row, 1)
        finally:
            shutil.rmtree(folder, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
