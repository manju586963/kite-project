#!/usr/bin/env python3
"""
Intraday Supertrend paper trading — historical test (Master Version).

Instrument : config.TRADING_SYMBOL (MCX CRUDEOIL August 2026 futures)
Timeframe  : 15-minute Heikin Ashi candles
Indicator  : Supertrend (10, 1)
Mode       : Intraday paper trading only — no Kite order APIs

Rules:
  - Convert raw OHLC → Heikin Ashi, then run Supertrend on HA candles
  - BUY  when Supertrend flips bearish → bullish (at HA candle close)
  - SELL when Supertrend flips bullish → bearish (at HA candle close)
  - BUY  closes SHORT then opens LONG
  - SELL closes LONG then opens SHORT
  - Max loss per trade: 125 points → square off (MAX_LOSS_STOP)
  - Stop checked before candle-close reversal signals
  - No new entries after 11:00 PM
  - Compulsory square-off at 11:15 PM (INTRADAY_SQUARE_OFF)
  - Each trading day starts FLAT (no overnight positions)

Outputs (never overwrite; dated folder in project root):
  01-aug-2026/
    supertrend_signals.csv
    paper_trades.csv
    paper_trading_summary.txt
    monthly/
      jun-2025.xlsx
      jul-2025.xlsx
      ...
    consolidated_monthly_summary.xlsx
    consolidated_monthly_summary.txt
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Iterable, Literal, Sequence

from dotenv import load_dotenv
from kiteconnect import KiteConnect
from openpyxl import Workbook

import config
from crude_oil_selector import resolve_contract

ROOT = Path(__file__).resolve().parent
ACCESS_TOKEN_FILE = ROOT / "access_token.txt"
SESSION_FILE = ROOT / ".kite_session.json"

INTERVAL = "15minute"
ST_PERIOD = 10
ST_MULTIPLIER = 1.0
DEFAULT_TRADING_DAYS = 30
# Extra calendar days so Supertrend has warm-up history before the test window.
DEFAULT_LOOKBACK_DAYS = 50
DEFAULT_MAX_LOSS_POINTS = 125.0

# Intraday session controls (IST clock on the candle timestamp).
NO_NEW_ENTRY_AFTER = time(23, 0)  # 11:00 PM — no new entries after this
SQUARE_OFF_AT = time(23, 15)  # 11:15 PM — compulsory flat

OUTPUT_ROOT = ROOT  # dated folders live in the project main folder
SIGNALS_NAME = "supertrend_signals.csv"
TRADES_NAME = "paper_trades.csv"
SUMMARY_NAME = "paper_trading_summary.txt"

Direction = Literal["bullish", "bearish"]
Side = Literal["BUY", "SELL"]


def make_output_dir(now: datetime | None = None) -> Path:
    """
    Create a dated subfolder in the project main folder. Never overwrite.

    Example: 01-aug-2026/
    If that folder already exists, use 01-aug-2026_2, _3, ...
    """
    stamp = (now or datetime.now()).strftime("%d-%b-%Y").lower()  # e.g. 01-aug-2026
    base = OUTPUT_ROOT / stamp
    path = base
    suffix = 2
    while path.exists():
        path = OUTPUT_ROOT / f"{stamp}_{suffix}"
        suffix += 1
    path.mkdir(parents=True, exist_ok=False)
    return path


def unique_file_path(directory: Path, filename: str) -> Path:
    """
    Return a path under directory that does not already exist.

    Example: supertrend_signals.csv → supertrend_signals_2.csv if needed.
    """
    candidate = directory / filename
    if not candidate.exists():
        return candidate

    stem = Path(filename).stem
    suffix = Path(filename).suffix
    n = 2
    while True:
        candidate = directory / f"{stem}_{n}{suffix}"
        if not candidate.exists():
            return candidate
        n += 1

# ---------------------------------------------------------------------------
# Credentials / Kite client
# ---------------------------------------------------------------------------
def load_api_key() -> str:
    load_dotenv(ROOT / ".env")
    api_key = os.getenv("KITE_API_KEY", "").strip()
    if api_key:
        return api_key
    if SESSION_FILE.exists():
        data = json.loads(SESSION_FILE.read_text(encoding="utf-8"))
        api_key = (data.get("api_key") or "").strip()
        if api_key:
            return api_key
    raise RuntimeError("KITE_API_KEY is missing from the .env file.")


def load_access_token() -> str:
    if ACCESS_TOKEN_FILE.exists():
        token = ACCESS_TOKEN_FILE.read_text(encoding="utf-8").strip()
        if token:
            return token
    if SESSION_FILE.exists():
        data = json.loads(SESSION_FILE.read_text(encoding="utf-8"))
        token = (data.get("access_token") or "").strip()
        if token:
            return token
    raise RuntimeError(
        "access_token.txt was not found. "
        "Complete today's Zerodha login first (python scripts/zerodha_login.py)."
    )


def make_kite() -> KiteConnect:
    kite = KiteConnect(api_key=load_api_key())
    kite.set_access_token(load_access_token())
    return kite


# ---------------------------------------------------------------------------
# Time helpers
# ---------------------------------------------------------------------------
def candle_clock(ts: datetime) -> time:
    """Return the candle's clock time (hour/minute as provided by Kite, IST)."""
    return time(ts.hour, ts.minute, ts.second)


def trading_day(ts: datetime) -> date:
    return ts.date()


def allows_new_entry(ts: datetime) -> bool:
    """Entries permitted at or before 11:00 PM candle close."""
    return candle_clock(ts) <= NO_NEW_ENTRY_AFTER


def is_square_off_candle(ts: datetime) -> bool:
    """Compulsory square-off at/after the 11:15 PM candle."""
    return candle_clock(ts) >= SQUARE_OFF_AT


# ---------------------------------------------------------------------------
# Supertrend (10, 1) — Wilder ATR
# ---------------------------------------------------------------------------
@dataclass
class Candle:
    date: datetime
    open: float
    high: float
    low: float
    close: float
    volume: float = 0.0
    oi: float | None = None


@dataclass
class CandleST(Candle):
    atr: float | None = None
    supertrend: float | None = None
    direction: Direction | None = None
    signal: Side | None = None


def to_heikin_ashi(candles: Sequence[Candle]) -> list[Candle]:
    """
    Convert standard OHLC candles to Heikin Ashi.

    HA_Close = (O + H + L + C) / 4
    HA_Open  = (prev_HA_Open + prev_HA_Close) / 2
               first bar uses (O + C) / 2
    HA_High  = max(H, HA_Open, HA_Close)
    HA_Low   = min(L, HA_Open, HA_Close)
    """
    if not candles:
        return []

    ha: list[Candle] = []
    for i, c in enumerate(candles):
        ha_close = (c.open + c.high + c.low + c.close) / 4.0
        if i == 0:
            ha_open = (c.open + c.close) / 2.0
        else:
            prev = ha[i - 1]
            ha_open = (prev.open + prev.close) / 2.0
        ha_high = max(c.high, ha_open, ha_close)
        ha_low = min(c.low, ha_open, ha_close)
        ha.append(
            Candle(
                date=c.date,
                open=ha_open,
                high=ha_high,
                low=ha_low,
                close=ha_close,
                volume=c.volume,
                oi=c.oi,
            )
        )
    return ha


def true_range(high: float, low: float, prev_close: float) -> float:
    return max(high - low, abs(high - prev_close), abs(low - prev_close))


def compute_supertrend(
    candles: Sequence[Candle],
    period: int = ST_PERIOD,
    multiplier: float = ST_MULTIPLIER,
) -> list[CandleST]:
    """
    Classic Supertrend using Wilder's ATR.

    Signals on completed candle close:
      BUY  when direction flips bearish → bullish
      SELL when direction flips bullish → bearish
    """
    if period < 1:
        raise ValueError("period must be >= 1")
    if not candles:
        return []

    n = len(candles)
    highs = [c.high for c in candles]
    lows = [c.low for c in candles]
    closes = [c.close for c in candles]

    tr = [0.0] * n
    atr: list[float | None] = [None] * n
    basic_ub = [0.0] * n
    basic_lb = [0.0] * n
    final_ub = [0.0] * n
    final_lb = [0.0] * n
    st: list[float | None] = [None] * n
    direction: list[Direction | None] = [None] * n

    tr[0] = highs[0] - lows[0]
    for i in range(1, n):
        tr[i] = true_range(highs[i], lows[i], closes[i - 1])

    if n >= period:
        seed = sum(tr[0:period]) / period
        atr[period - 1] = seed
        for i in range(period, n):
            prev = atr[i - 1]
            assert prev is not None
            atr[i] = ((prev * (period - 1)) + tr[i]) / period

    for i in range(n):
        if atr[i] is None:
            continue
        mid = (highs[i] + lows[i]) / 2.0
        basic_ub[i] = mid + multiplier * atr[i]
        basic_lb[i] = mid - multiplier * atr[i]

        if i == period - 1:
            final_ub[i] = basic_ub[i]
            final_lb[i] = basic_lb[i]
        else:
            prev_fub = final_ub[i - 1]
            prev_flb = final_lb[i - 1]
            final_ub[i] = (
                basic_ub[i]
                if (basic_ub[i] < prev_fub or closes[i - 1] > prev_fub)
                else prev_fub
            )
            final_lb[i] = (
                basic_lb[i]
                if (basic_lb[i] > prev_flb or closes[i - 1] < prev_flb)
                else prev_flb
            )

        if i == period - 1:
            if closes[i] <= final_ub[i]:
                st[i] = final_ub[i]
                direction[i] = "bearish"
            else:
                st[i] = final_lb[i]
                direction[i] = "bullish"
            continue

        prev_dir = direction[i - 1]
        if prev_dir is None:
            continue

        if prev_dir == "bullish":
            if closes[i] < final_lb[i]:
                st[i] = final_ub[i]
                direction[i] = "bearish"
            else:
                st[i] = final_lb[i]
                direction[i] = "bullish"
        else:
            if closes[i] > final_ub[i]:
                st[i] = final_lb[i]
                direction[i] = "bullish"
            else:
                st[i] = final_ub[i]
                direction[i] = "bearish"

    rows: list[CandleST] = []
    for i, c in enumerate(candles):
        signal: Side | None = None
        if (
            i > 0
            and direction[i]
            and direction[i - 1]
            and direction[i] != direction[i - 1]
        ):
            signal = "BUY" if direction[i] == "bullish" else "SELL"
        rows.append(
            CandleST(
                date=c.date,
                open=c.open,
                high=c.high,
                low=c.low,
                close=c.close,
                volume=c.volume,
                oi=c.oi,
                atr=atr[i],
                supertrend=st[i],
                direction=direction[i],
                signal=signal,
            )
        )
    return rows


# ---------------------------------------------------------------------------
# Intraday paper trading engine
# ---------------------------------------------------------------------------
@dataclass
class PaperTrade:
    trade_id: int
    side: Side  # BUY=LONG, SELL=SHORT
    entry_time: datetime
    entry_price: float
    trade_date: date | None = None
    stop_price: float | None = None
    exit_time: datetime | None = None
    exit_price: float | None = None
    points: float | None = None
    pnl: float | None = None
    exit_reason: str | None = None


def stop_price_for(side: Side, entry_price: float, max_loss_points: float) -> float:
    """LONG: entry − N points. SHORT: entry + N points."""
    if side == "BUY":
        return entry_price - max_loss_points
    return entry_price + max_loss_points


def hit_max_loss_stop(
    trade: PaperTrade,
    candle: CandleST,
    max_loss_points: float,
) -> float | None:
    """
    If candle high/low touches the stop, return the exact stop fill price.
    Gaps/slippage are excluded (fills at stop).
    """
    stop = trade.stop_price
    if stop is None:
        stop = stop_price_for(trade.side, trade.entry_price, max_loss_points)
    if trade.side == "BUY":
        if float(candle.low) <= stop:
            return stop
    else:
        if float(candle.high) >= stop:
            return stop
    return None


def simulate_paper_trades(
    rows: Sequence[CandleST],
    *,
    lot_size: int,
    max_loss_points: float = DEFAULT_MAX_LOSS_POINTS,
) -> list[PaperTrade]:
    """
    Intraday-only paper engine with max-loss square-off.

    Per candle:
      1) Day change → prior day INTRADAY_SQUARE_OFF if still open
      2) Max-loss stop on high/low (before reversal signal) → MAX_LOSS_STOP
      3) 11:15 PM compulsory square-off
      4) Supertrend reversal at close (may open a new position)

    Stop:
      LONG  = entry − max_loss_points
      SHORT = entry + max_loss_points
      Paper loss capped at max_loss_points × lot_size
    """
    if max_loss_points <= 0:
        raise ValueError("max_loss_points must be > 0")

    trades: list[PaperTrade] = []
    open_trade: PaperTrade | None = None
    next_id = 1
    current_day: date | None = None
    last_row: CandleST | None = None

    def close_open(exit_time: datetime, exit_price: float, reason: str) -> None:
        nonlocal open_trade
        if open_trade is None:
            return
        if open_trade.side == "BUY":
            points = exit_price - open_trade.entry_price
        else:
            points = open_trade.entry_price - exit_price
        open_trade.exit_time = exit_time
        open_trade.exit_price = exit_price
        open_trade.points = points
        open_trade.pnl = points * lot_size
        open_trade.exit_reason = reason
        trades.append(open_trade)
        open_trade = None

    def open_position(side: Side, ts: datetime, price: float, day: date) -> None:
        nonlocal open_trade, next_id
        open_trade = PaperTrade(
            trade_id=next_id,
            side=side,
            entry_time=ts,
            entry_price=price,
            trade_date=day,
            stop_price=stop_price_for(side, price, max_loss_points),
        )
        next_id += 1

    for row in rows:
        day = trading_day(row.date)

        # New calendar trading day → force flat using previous day's last candle.
        if current_day is None:
            current_day = day
        elif day != current_day:
            if open_trade is not None and last_row is not None:
                close_open(
                    last_row.date,
                    float(last_row.close),
                    "INTRADAY_SQUARE_OFF",
                )
            current_day = day

        # Max-loss stop BEFORE reversal / session square-off close.
        # Skip the entry candle (position opens at that candle's close).
        if open_trade is not None and row.date != open_trade.entry_time:
            stop_exit = hit_max_loss_stop(open_trade, row, max_loss_points)
            if stop_exit is not None:
                close_open(row.date, stop_exit, "MAX_LOSS_STOP")

        # Compulsory square-off candle (11:15 PM+).
        if is_square_off_candle(row.date):
            if open_trade is not None:
                close_open(row.date, float(row.close), "INTRADAY_SQUARE_OFF")
            last_row = row
            continue

        # Supertrend flips only inside the entry window.
        if row.signal and allows_new_entry(row.date):
            price = float(row.close)
            ts = row.date

            if open_trade is None:
                open_position(row.signal, ts, price, day)
            elif row.signal != open_trade.side:
                close_open(ts, price, f"FLIP_{row.signal}")
                open_position(row.signal, ts, price, day)

        last_row = row

    if open_trade is not None and last_row is not None:
        close_open(last_row.date, float(last_row.close), "INTRADAY_SQUARE_OFF")

    return trades


def filter_last_trading_days(
    rows: Sequence[CandleST],
    trading_days: int = DEFAULT_TRADING_DAYS,
) -> list[CandleST]:
    """Keep candles belonging to the last N unique trading dates."""
    if trading_days <= 0 or not rows:
        return list(rows)
    days = sorted({trading_day(r.date) for r in rows})
    keep = set(days[-trading_days:])
    return [r for r in rows if trading_day(r.date) in keep]


def summarize_trades(trades: Sequence[PaperTrade]) -> dict:
    closed = [t for t in trades if t.pnl is not None]
    wins = [t for t in closed if (t.pnl or 0) > 0]
    losses = [t for t in closed if (t.pnl or 0) < 0]
    flats = [t for t in closed if (t.pnl or 0) == 0]
    square_offs = [t for t in closed if t.exit_reason == "INTRADAY_SQUARE_OFF"]
    max_loss_stops = [t for t in closed if t.exit_reason == "MAX_LOSS_STOP"]
    net = sum(t.pnl or 0 for t in closed)

    by_day: dict[date, float] = {}
    for t in closed:
        d = t.trade_date or trading_day(t.entry_time)
        by_day[d] = by_day.get(d, 0.0) + (t.pnl or 0.0)

    return {
        "total_trades": len(closed),
        "wins": len(wins),
        "losses": len(losses),
        "flats": len(flats),
        "square_offs": len(square_offs),
        "max_loss_stops": len(max_loss_stops),
        "trading_days": len(by_day),
        "win_rate_pct": (len(wins) / len(closed) * 100.0) if closed else 0.0,
        "net_pnl": net,
        "gross_profit": sum(t.pnl or 0 for t in wins),
        "gross_loss": sum(t.pnl or 0 for t in losses),
        "daily_pnl": dict(sorted(by_day.items())),
    }


# ---------------------------------------------------------------------------
# Historical download
# ---------------------------------------------------------------------------
def candles_from_kite_records(records: Iterable[dict]) -> list[Candle]:
    out: list[Candle] = []
    for r in records:
        out.append(
            Candle(
                date=(
                    r["date"]
                    if isinstance(r["date"], datetime)
                    else datetime.fromisoformat(str(r["date"]))
                ),
                open=float(r["open"]),
                high=float(r["high"]),
                low=float(r["low"]),
                close=float(r["close"]),
                volume=float(r.get("volume") or 0),
                oi=float(r["oi"]) if r.get("oi") is not None else None,
            )
        )
    out.sort(key=lambda c: c.date)
    return out


def fetch_historical_15m(
    kite: KiteConnect,
    instrument_token: int,
    from_dt: datetime,
    to_dt: datetime,
) -> list[Candle]:
    """Download 15-minute candles in ~60-day chunks and merge."""
    chunk_days = 60
    cursor = from_dt
    all_records: list[dict] = []

    while cursor < to_dt:
        chunk_end = min(cursor + timedelta(days=chunk_days), to_dt)
        print(
            f"  Fetching {cursor:%Y-%m-%d} → {chunk_end:%Y-%m-%d} ({INTERVAL})..."
        )
        records = kite.historical_data(
            instrument_token,
            cursor,
            chunk_end,
            INTERVAL,
            continuous=False,
            oi=True,
        )
        all_records.extend(records)
        cursor = chunk_end + timedelta(seconds=1)

    by_ts: dict[datetime, dict] = {}
    for r in all_records:
        by_ts[r["date"]] = r
    ordered = [by_ts[k] for k in sorted(by_ts)]
    return candles_from_kite_records(ordered)


# ---------------------------------------------------------------------------
# Output writers
# ---------------------------------------------------------------------------
def write_signals_csv(path: Path, rows: Sequence[CandleST], symbol: str) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "datetime",
                "symbol",
                "ha_open",
                "ha_high",
                "ha_low",
                "ha_close",
                "volume",
                "atr",
                "supertrend",
                "direction",
                "signal",
                "entry_allowed",
            ]
        )
        for r in rows:
            writer.writerow(
                [
                    r.date.isoformat(),
                    symbol,
                    r.open,
                    r.high,
                    r.low,
                    r.close,
                    r.volume,
                    "" if r.atr is None else round(r.atr, 4),
                    "" if r.supertrend is None else round(r.supertrend, 4),
                    r.direction or "",
                    r.signal or "",
                    (
                        "YES"
                        if allows_new_entry(r.date) and not is_square_off_candle(r.date)
                        else "NO"
                    ),
                ]
            )


def write_trades_csv(
    path: Path, trades: Sequence[PaperTrade], symbol: str, lot_size: int
) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "trade_id",
                "trade_date",
                "symbol",
                "side",
                "position",
                "entry_time",
                "entry_price",
                "stop_price",
                "exit_time",
                "exit_price",
                "points",
                "lot_size",
                "pnl",
                "exit_reason",
            ]
        )
        for t in trades:
            writer.writerow(
                [
                    t.trade_id,
                    (t.trade_date or trading_day(t.entry_time)).isoformat(),
                    symbol,
                    t.side,
                    "LONG" if t.side == "BUY" else "SHORT",
                    t.entry_time.isoformat(),
                    t.entry_price,
                    "" if t.stop_price is None else t.stop_price,
                    "" if t.exit_time is None else t.exit_time.isoformat(),
                    "" if t.exit_price is None else t.exit_price,
                    "" if t.points is None else round(t.points, 4),
                    lot_size,
                    "" if t.pnl is None else round(t.pnl, 4),
                    t.exit_reason or "",
                ]
            )


def _month_key(ts: datetime | date) -> str:
    """Return folder/file month stamp like 2025-06."""
    if isinstance(ts, datetime):
        return f"{ts.year:04d}-{ts.month:02d}"
    return f"{ts.year:04d}-{ts.month:02d}"


def _month_label(year: int, month: int) -> str:
    """Human month name for Excel, e.g. jun-2025."""
    return datetime(year, month, 1).strftime("%b-%Y").lower()


def write_monthly_excel(
    output_dir: Path,
    rows: Sequence[CandleST],
    trades: Sequence[PaperTrade],
    *,
    symbol: str,
    lot_size: int,
    max_loss_points: float,
) -> list[Path]:
    """
    Write one Excel workbook per calendar month under output_dir/monthly/.

    Each file contains sheets: signals, trades, summary.
    Existing files are never overwritten (unique_file_path).
    """
    monthly_dir = output_dir / "monthly"
    monthly_dir.mkdir(parents=True, exist_ok=True)

    rows_by_month: dict[str, list[CandleST]] = {}
    for r in rows:
        rows_by_month.setdefault(_month_key(r.date), []).append(r)

    trades_by_month: dict[str, list[PaperTrade]] = {}
    for t in trades:
        d = t.trade_date or trading_day(t.entry_time)
        trades_by_month.setdefault(_month_key(d), []).append(t)

    months = sorted(set(rows_by_month) | set(trades_by_month))
    written: list[Path] = []

    for month in months:
        year_s, month_s = month.split("-")
        year, mon = int(year_s), int(month_s)
        label = _month_label(year, mon)
        path = unique_file_path(monthly_dir, f"{label}.xlsx")

        month_rows = rows_by_month.get(month, [])
        month_trades = trades_by_month.get(month, [])
        month_stats = summarize_trades(month_trades)
        month_signals = [
            r for r in month_rows if r.signal and allows_new_entry(r.date)
        ]

        wb = Workbook()

        # --- signals ---
        ws_sig = wb.active
        ws_sig.title = "signals"
        sig_headers = [
            "datetime",
            "symbol",
            "ha_open",
            "ha_high",
            "ha_low",
            "ha_close",
            "volume",
            "atr",
            "supertrend",
            "direction",
            "signal",
            "entry_allowed",
        ]
        ws_sig.append(sig_headers)
        for r in month_rows:
            ws_sig.append(
                [
                    r.date.isoformat(),
                    symbol,
                    r.open,
                    r.high,
                    r.low,
                    r.close,
                    r.volume,
                    None if r.atr is None else round(r.atr, 4),
                    None if r.supertrend is None else round(r.supertrend, 4),
                    r.direction or "",
                    r.signal or "",
                    (
                        "YES"
                        if allows_new_entry(r.date) and not is_square_off_candle(r.date)
                        else "NO"
                    ),
                ]
            )

        # --- trades ---
        ws_tr = wb.create_sheet("trades")
        tr_headers = [
            "trade_id",
            "trade_date",
            "symbol",
            "side",
            "position",
            "entry_time",
            "entry_price",
            "stop_price",
            "exit_time",
            "exit_price",
            "points",
            "lot_size",
            "pnl",
            "exit_reason",
        ]
        ws_tr.append(tr_headers)
        for t in month_trades:
            ws_tr.append(
                [
                    t.trade_id,
                    (t.trade_date or trading_day(t.entry_time)).isoformat(),
                    symbol,
                    t.side,
                    "LONG" if t.side == "BUY" else "SHORT",
                    t.entry_time.isoformat(),
                    t.entry_price,
                    t.stop_price,
                    None if t.exit_time is None else t.exit_time.isoformat(),
                    t.exit_price,
                    None if t.points is None else round(t.points, 4),
                    lot_size,
                    None if t.pnl is None else round(t.pnl, 4),
                    t.exit_reason or "",
                ]
            )

        # --- summary ---
        ws_sum = wb.create_sheet("summary")
        summary_rows = [
            ("Month", label),
            ("Symbol", symbol),
            ("Interval", f"{INTERVAL} (Heikin Ashi)"),
            ("Supertrend", f"({ST_PERIOD}, {ST_MULTIPLIER:g})"),
            ("Max loss points", max_loss_points),
            ("Lot size", lot_size),
            ("Candles", len(month_rows)),
            ("Signals", len(month_signals)),
            ("Total trades", month_stats["total_trades"]),
            ("Wins", month_stats["wins"]),
            ("Losses", month_stats["losses"]),
            ("Max-loss stops", month_stats["max_loss_stops"]),
            ("Square-offs", month_stats["square_offs"]),
            ("Win rate %", round(month_stats["win_rate_pct"], 2)),
            ("Gross profit", round(month_stats["gross_profit"], 2)),
            ("Gross loss", round(month_stats["gross_loss"], 2)),
            ("Net paper P&L", round(month_stats["net_pnl"], 2)),
        ]
        ws_sum.append(["Field", "Value"])
        for field, value in summary_rows:
            ws_sum.append([field, value])

        if month_stats.get("daily_pnl"):
            ws_sum.append([])
            ws_sum.append(["Daily realised paper P&L", ""])
            ws_sum.append(["Date", "PnL"])
            for d, pnl in month_stats["daily_pnl"].items():
                ws_sum.append([d.isoformat(), round(pnl, 2)])

        wb.save(path)
        written.append(path)

    return written


def write_consolidated_monthly_summary(
    output_dir: Path,
    rows: Sequence[CandleST],
    trades: Sequence[PaperTrade],
    *,
    symbol: str,
    lot_size: int,
    max_loss_points: float,
    overall_stats: dict,
) -> tuple[Path, Path]:
    """
    Consolidate all months into one summary Excel + one text file.

    Excel sheets:
      - monthly_summary : one row per month + TOTAL
      - all_trades      : every paper trade
      - daily_pnl       : daily realised P&L across the run
    """
    rows_by_month: dict[str, list[CandleST]] = {}
    for r in rows:
        rows_by_month.setdefault(_month_key(r.date), []).append(r)

    trades_by_month: dict[str, list[PaperTrade]] = {}
    for t in trades:
        d = t.trade_date or trading_day(t.entry_time)
        trades_by_month.setdefault(_month_key(d), []).append(t)

    months = sorted(set(rows_by_month) | set(trades_by_month))

    xlsx_path = unique_file_path(output_dir, "consolidated_monthly_summary.xlsx")
    txt_path = unique_file_path(output_dir, "consolidated_monthly_summary.txt")

    wb = Workbook()
    ws = wb.active
    ws.title = "monthly_summary"
    headers = [
        "month",
        "candles",
        "signals",
        "trades",
        "wins",
        "losses",
        "flats",
        "max_loss_stops",
        "square_offs",
        "win_rate_pct",
        "gross_profit",
        "gross_loss",
        "net_pnl",
    ]
    ws.append(headers)

    txt_lines = [
        "Consolidated Monthly Paper Trading Summary",
        "==========================================",
        "",
        f"Symbol           : {symbol}",
        f"Supertrend       : ({ST_PERIOD}, {ST_MULTIPLIER:g}) on Heikin Ashi",
        f"Max loss points  : {max_loss_points:g}",
        f"Lot size         : {lot_size}",
        f"Overall trades   : {overall_stats['total_trades']}",
        f"Overall win rate : {overall_stats['win_rate_pct']:.2f}%",
        f"Overall net P&L  : {overall_stats['net_pnl']:.2f}",
        "",
        f"{'Month':<12} {'Trades':>7} {'Wins':>6} {'Losses':>7} "
        f"{'Stops':>6} {'SqOff':>6} {'Win%':>8} {'Net P&L':>12}",
        "-" * 72,
    ]

    totals = {
        "candles": 0,
        "signals": 0,
        "trades": 0,
        "wins": 0,
        "losses": 0,
        "flats": 0,
        "max_loss_stops": 0,
        "square_offs": 0,
        "gross_profit": 0.0,
        "gross_loss": 0.0,
        "net_pnl": 0.0,
    }

    for month in months:
        year_s, month_s = month.split("-")
        label = _month_label(int(year_s), int(month_s))
        month_rows = rows_by_month.get(month, [])
        month_trades = trades_by_month.get(month, [])
        st = summarize_trades(month_trades)
        sig_count = len(
            [r for r in month_rows if r.signal and allows_new_entry(r.date)]
        )

        row_vals = [
            label,
            len(month_rows),
            sig_count,
            st["total_trades"],
            st["wins"],
            st["losses"],
            st["flats"],
            st["max_loss_stops"],
            st["square_offs"],
            round(st["win_rate_pct"], 2),
            round(st["gross_profit"], 2),
            round(st["gross_loss"], 2),
            round(st["net_pnl"], 2),
        ]
        ws.append(row_vals)

        totals["candles"] += len(month_rows)
        totals["signals"] += sig_count
        totals["trades"] += st["total_trades"]
        totals["wins"] += st["wins"]
        totals["losses"] += st["losses"]
        totals["flats"] += st["flats"]
        totals["max_loss_stops"] += st["max_loss_stops"]
        totals["square_offs"] += st["square_offs"]
        totals["gross_profit"] += st["gross_profit"]
        totals["gross_loss"] += st["gross_loss"]
        totals["net_pnl"] += st["net_pnl"]

        txt_lines.append(
            f"{label:<12} {st['total_trades']:>7} {st['wins']:>6} {st['losses']:>7} "
            f"{st['max_loss_stops']:>6} {st['square_offs']:>6} "
            f"{st['win_rate_pct']:>7.2f}% {st['net_pnl']:>12.2f}"
        )

    overall_win = (
        (totals["wins"] / totals["trades"] * 100.0) if totals["trades"] else 0.0
    )
    ws.append(
        [
            "TOTAL",
            totals["candles"],
            totals["signals"],
            totals["trades"],
            totals["wins"],
            totals["losses"],
            totals["flats"],
            totals["max_loss_stops"],
            totals["square_offs"],
            round(overall_win, 2),
            round(totals["gross_profit"], 2),
            round(totals["gross_loss"], 2),
            round(totals["net_pnl"], 2),
        ]
    )
    txt_lines.append("-" * 72)
    txt_lines.append(
        f"{'TOTAL':<12} {totals['trades']:>7} {totals['wins']:>6} {totals['losses']:>7} "
        f"{totals['max_loss_stops']:>6} {totals['square_offs']:>6} "
        f"{overall_win:>7.2f}% {totals['net_pnl']:>12.2f}"
    )
    txt_lines.append("")

    # all_trades sheet
    ws_tr = wb.create_sheet("all_trades")
    ws_tr.append(
        [
            "month",
            "trade_id",
            "trade_date",
            "symbol",
            "side",
            "position",
            "entry_time",
            "entry_price",
            "stop_price",
            "exit_time",
            "exit_price",
            "points",
            "lot_size",
            "pnl",
            "exit_reason",
        ]
    )
    for t in trades:
        d = t.trade_date or trading_day(t.entry_time)
        label = _month_label(d.year, d.month)
        ws_tr.append(
            [
                label,
                t.trade_id,
                d.isoformat(),
                symbol,
                t.side,
                "LONG" if t.side == "BUY" else "SHORT",
                t.entry_time.isoformat(),
                t.entry_price,
                t.stop_price,
                None if t.exit_time is None else t.exit_time.isoformat(),
                t.exit_price,
                None if t.points is None else round(t.points, 4),
                lot_size,
                None if t.pnl is None else round(t.pnl, 4),
                t.exit_reason or "",
            ]
        )

    # daily_pnl sheet
    ws_day = wb.create_sheet("daily_pnl")
    ws_day.append(["month", "date", "pnl"])
    daily = overall_stats.get("daily_pnl") or {}
    for d, pnl in daily.items():
        ws_day.append([_month_label(d.year, d.month), d.isoformat(), round(pnl, 2)])

    wb.save(xlsx_path)
    txt_path.write_text("\n".join(txt_lines) + "\n", encoding="utf-8")
    return xlsx_path, txt_path


def write_summary(
    path: Path,
    *,
    symbol: str,
    token: int,
    lot_size: int,
    from_dt: datetime,
    to_dt: datetime,
    candle_count: int,
    signal_count: int,
    trades: Sequence[PaperTrade],
    stats: dict,
    max_loss_points: float = DEFAULT_MAX_LOSS_POINTS,
) -> None:
    lines = [
        "Intraday Supertrend Paper Trading Summary",
        "=========================================",
        "",
        f"Symbol              : {symbol}",
        f"Instrument token    : {token}",
        f"Interval            : {INTERVAL} (Heikin Ashi)",
        f"Supertrend          : ({ST_PERIOD}, {ST_MULTIPLIER:g}) on Heikin Ashi",
        f"Style               : Intraday (daily square-off)",
        f"Max loss / trade    : {max_loss_points:g} points "
        f"(max paper loss {max_loss_points:g} × lot_size)",
        f"No new entries after: {NO_NEW_ENTRY_AFTER.strftime('%I:%M %p')}",
        f"Compulsory exit     : {SQUARE_OFF_AT.strftime('%I:%M %p')} (INTRADAY_SQUARE_OFF)",
        f"Lot size            : {lot_size}",
        f"From                : {from_dt.isoformat()}",
        f"To                  : {to_dt.isoformat()}",
        f"Candles             : {candle_count}",
        f"Signals (BUY/SELL)  : {signal_count}",
        f"Trading days        : {stats['trading_days']}",
        "",
        f"Total trades        : {stats['total_trades']}",
        f"Wins                : {stats['wins']}",
        f"Losses              : {stats['losses']}",
        f"Flats               : {stats['flats']}",
        f"Max-loss stops      : {stats['max_loss_stops']}",
        f"Square-offs         : {stats['square_offs']}",
        f"Win rate            : {stats['win_rate_pct']:.2f}%",
        f"Gross profit        : {stats['gross_profit']:.2f}",
        f"Gross loss          : {stats['gross_loss']:.2f}",
        f"Net paper P&L       : {stats['net_pnl']:.2f}",
        "",
        "Notes:",
        "- Paper trading only. No real orders were placed.",
        "- Candles are Heikin Ashi conversions of Kite 15-minute OHLC.",
        "- Entry uses HA signal candle close; MAX_LOSS_STOP fills at the HA stop price.",
        "- LONG stop = entry − max-loss points; SHORT stop = entry + max-loss points.",
        "- Stop-loss is checked before candle-close reversal signals.",
        "- Each day starts FLAT; overnight positions are not permitted.",
        "- Forced session exits are recorded as INTRADAY_SQUARE_OFF.",
        "- P&L = points × lot_size (from the instrument master).",
        "",
    ]
    if stats.get("daily_pnl"):
        lines.append("Daily realised paper P&L:")
        for d, pnl in stats["daily_pnl"].items():
            lines.append(f"  {d.isoformat()} : {pnl:.2f}")
        lines.append("")
    if trades:
        lines.append("Trade log (brief):")
        for t in trades:
            lines.append(
                f"  #{t.trade_id} {t.side:4} "
                f"{t.entry_price} → {t.exit_price} "
                f"pts={t.points} pnl={t.pnl} ({t.exit_reason})"
            )
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Intraday Supertrend (10,1) on Heikin Ashi 15-minute candles"
        )
    )
    parser.add_argument(
        "--days",
        type=int,
        default=DEFAULT_LOOKBACK_DAYS,
        help=f"Calendar lookback days to download (default {DEFAULT_LOOKBACK_DAYS})",
    )
    parser.add_argument(
        "--trading-days",
        type=int,
        default=DEFAULT_TRADING_DAYS,
        help=f"Use the last N trading days for the test (default {DEFAULT_TRADING_DAYS})",
    )
    parser.add_argument("--from", dest="from_date", help="Start date YYYY-MM-DD")
    parser.add_argument("--to", dest="to_date", help="End date YYYY-MM-DD")
    parser.add_argument(
        "--csv-candles",
        help="Optional local candles CSV (datetime,open,high,low,close[,volume]) "
        "to run without calling Kite historical API",
    )
    parser.add_argument(
        "--max-loss-points",
        type=float,
        default=DEFAULT_MAX_LOSS_POINTS,
        help=f"Max loss in points per trade before square-off "
        f"(default {DEFAULT_MAX_LOSS_POINTS:g})",
    )
    return parser.parse_args(argv)


def load_candles_csv(path: Path) -> list[Candle]:
    rows: list[Candle] = []
    with path.open(encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        for r in reader:
            ts = r.get("datetime") or r.get("date")
            if not ts:
                continue
            rows.append(
                Candle(
                    date=datetime.fromisoformat(ts.replace("Z", "+00:00")),
                    open=float(r["open"]),
                    high=float(r["high"]),
                    low=float(r["low"]),
                    close=float(r["close"]),
                    volume=float(r["volume"]) if r.get("volume") else 0.0,
                )
            )
    rows.sort(key=lambda c: c.date)
    return rows


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)

    symbol = config.TRADING_SYMBOL
    print("========================================")
    print("INTRADAY SUPERTREND PAPER TRADING")
    print("========================================")
    print(f"Symbol       : {symbol}")
    print(f"Interval     : {INTERVAL} (Heikin Ashi)")
    print(f"Supertrend   : ({ST_PERIOD}, {ST_MULTIPLIER:g}) on HA")
    print(f"No entries   : after {NO_NEW_ENTRY_AFTER.strftime('%I:%M %p')}")
    print(f"Square-off   : {SQUARE_OFF_AT.strftime('%I:%M %p')}")
    print(f"Max loss     : {args.max_loss_points:g} points / trade")
    print("Overnight    : NOT PERMITTED")
    print("Orders       : DISABLED (paper only)")
    print("========================================")
    print()

    if args.csv_candles:
        print(f"Loading candles from {args.csv_candles}...")
        candles = load_candles_csv(Path(args.csv_candles))
        try:
            contract = resolve_contract()
            token = contract.instrument_token
            lot_size = int(contract.lot_size)
            symbol = contract.tradingsymbol
        except Exception:
            token = 0
            lot_size = 1
        from_dt = candles[0].date if candles else datetime.now()
        to_dt = candles[-1].date if candles else datetime.now()
    else:
        print(f"Resolving contract {symbol!r}...")
        contract = resolve_contract()
        token = int(contract.instrument_token)
        lot_size = int(contract.lot_size)
        symbol = contract.tradingsymbol
        print(f"  token={token}  expiry={contract.expiry}  lot_size={lot_size}")

        to_dt = (
            datetime.strptime(args.to_date, "%Y-%m-%d")
            if args.to_date
            else datetime.now()
        )
        from_dt = (
            datetime.strptime(args.from_date, "%Y-%m-%d")
            if args.from_date
            else to_dt - timedelta(days=args.days)
        )

        kite = make_kite()
        print("Downloading historical 15-minute candles...")
        candles = fetch_historical_15m(kite, token, from_dt, to_dt)

    if len(candles) < ST_PERIOD + 2:
        raise RuntimeError(
            f"Not enough candles ({len(candles)}) to compute Supertrend({ST_PERIOD})."
        )

    print(f"Candles loaded: {len(candles)} (raw OHLC)")
    print("Converting to Heikin Ashi...")
    ha_candles = to_heikin_ashi(candles)
    print("Computing Supertrend (10,1) on Heikin Ashi...")
    all_rows = compute_supertrend(ha_candles, ST_PERIOD, ST_MULTIPLIER)

    # Supertrend uses full history for warm-up; simulate only the test window.
    rows = filter_last_trading_days(all_rows, args.trading_days)
    test_days = sorted({trading_day(r.date) for r in rows})
    print(
        f"Test window   : {test_days[0]} → {test_days[-1]} "
        f"({len(test_days)} trading days, {len(rows)} candles)"
    )

    print("Simulating intraday paper trades...")
    signals = [r for r in rows if r.signal and allows_new_entry(r.date)]
    trades = simulate_paper_trades(
        rows,
        lot_size=lot_size,
        max_loss_points=args.max_loss_points,
    )
    stats = summarize_trades(trades)

    output_dir = make_output_dir()
    signals_file = unique_file_path(output_dir, SIGNALS_NAME)
    trades_file = unique_file_path(output_dir, TRADES_NAME)
    summary_file = unique_file_path(output_dir, SUMMARY_NAME)

    write_signals_csv(signals_file, rows, symbol)
    write_trades_csv(trades_file, trades, symbol, lot_size)
    write_summary(
        summary_file,
        symbol=symbol,
        token=token,
        lot_size=lot_size,
        from_dt=rows[0].date if rows else from_dt,
        to_dt=rows[-1].date if rows else to_dt,
        candle_count=len(rows),
        signal_count=len(signals),
        trades=trades,
        stats=stats,
        max_loss_points=args.max_loss_points,
    )
    excel_files = write_monthly_excel(
        output_dir,
        rows,
        trades,
        symbol=symbol,
        lot_size=lot_size,
        max_loss_points=args.max_loss_points,
    )
    consolidated_xlsx, consolidated_txt = write_consolidated_monthly_summary(
        output_dir,
        rows,
        trades,
        symbol=symbol,
        lot_size=lot_size,
        max_loss_points=args.max_loss_points,
        overall_stats=stats,
    )

    print()
    print("Done (intraday paper trading only — no real orders).")
    print(f"  Output folder: {output_dir}")
    print(f"  Signals file : {signals_file.name}")
    print(f"  Trades file  : {trades_file.name} ({stats['total_trades']} trades)")
    print(f"  Summary file : {summary_file.name}")
    print(f"  Monthly Excel: {len(excel_files)} file(s) in {output_dir / 'monthly'}")
    for xf in excel_files:
        print(f"    - {xf.relative_to(output_dir)}")
    print(f"  Consolidated : {consolidated_xlsx.name}")
    print(f"  Consolidated : {consolidated_txt.name}")
    print(f"  Max-loss stops: {stats['max_loss_stops']}")
    print(f"  Square-offs  : {stats['square_offs']}")
    print(f"  Win rate     : {stats['win_rate_pct']:.2f}%")
    print(f"  Net paper P&L: {stats['net_pnl']:.2f}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001
        print(f"Historical paper trading failed: {exc}", file=sys.stderr)
        raise SystemExit(1)
